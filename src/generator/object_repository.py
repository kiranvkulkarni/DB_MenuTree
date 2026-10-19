"""The object repository: every UI element a run met, with all its locators.

A central store of locators kept apart from the tests that use them
(agreed 2026-10-20). One entry per element, under a stable name built from
its screen and its label -- `PORTRAIT.QuickControls.Flash` -- holding every
way to find it:

    text, content_desc, resource_id, class, xpath, uiautomator

and the one the walk itself used (`preferred`, in the confirmed priority
text > description > resource id > xpath). Written three ways per run:

    <run>_object_repository.json   for any tool
    sheet "ObjectRepository"        in the MenuTree workbook
    <run>_page_objects.py           one class per screen, for test code

The MenuTree rows and the UVTA cases name their element, so a test can be
moved onto the repository without re-walking anything.
"""
import json
import keyword
import re
from collections import OrderedDict
from datetime import datetime
from pathlib import Path
from typing import Dict, List, Optional, Sequence, Tuple

LONG_PRESS = " (long press)"


def _words(text: str) -> List[str]:
    return re.findall(r"[A-Za-z0-9]+", text or "")


def _camel(text: str) -> str:
    """'Quick controls' -> 'QuickControls'; 'PORTRAIT' stays as written."""
    parts = _words(text)
    if not parts:
        return "Item"
    out = "".join(p if p.isupper() and len(p) > 1 else p[:1].upper() + p[1:] for p in parts)
    return out[:60]


def _segment(label: str) -> str:
    long = label.endswith(LONG_PRESS)
    base = label[:-len(LONG_PRESS)] if long else label
    return _camel(base) + ("LongPress" if long else "")


def screen_name(path: Sequence[str]) -> str:
    return ".".join(_segment(p) for p in path) if path else "Main"


def _uiautomator(by: str, value: str, resource_id: str) -> str:
    def q(v):
        return '"' + str(v).replace("\\", "\\\\").replace('"', '\\"') + '"'
    if by == "text":
        return f"new UiSelector().text({q(value)})"
    if by == "desc":
        return f"new UiSelector().description({q(value)})"
    if by == "id":
        return f"new UiSelector().resourceId({q(resource_id or value)})"
    return ""


def build(rows: Sequence[Dict]) -> Tuple["OrderedDict[str, Dict]", List[str]]:
    """(entries by name, the entry name for each row -- "" when none)."""
    entries: "OrderedDict[str, Dict]" = OrderedDict()
    by_identity: Dict[tuple, str] = {}
    names: List[str] = []
    for row in rows:
        loc = row.get("locators") or {}
        if not loc or not (loc.get("text") or loc.get("content_desc")
                           or loc.get("resource_id") or loc.get("xpath")):
            names.append("")
            continue
        screen = list(row.get("path") or [])
        raw = row.get("raw_label") or row.get("label") or ""
        if (row.get("label") or "").endswith(LONG_PRESS):
            raw += LONG_PRESS
        identity = (tuple(screen), raw, loc.get("resource_id"), loc.get("xpath"))
        if identity in by_identity:
            names.append(by_identity[identity])
            continue
        base = f"{screen_name(screen)}.{_segment(raw)}"
        name, n = base, 2
        while name in entries:
            name, n = f"{base}_{n}", n + 1
        by = row.get("selector_kind") or ""
        value = row.get("selector_value") or ""
        entries[name] = {
            "name": name,
            "screen": " > ".join(screen) or "(main screen)",
            "label": raw,
            "kind": row.get("kind", ""),
            "long_press": (row.get("label") or "").endswith(LONG_PRESS),
            "package": loc.get("package", ""),
            "preferred": {"by": by, "value": value},
            "locators": {
                "text": loc.get("text", ""),
                "content_desc": loc.get("content_desc", ""),
                "resource_id": loc.get("resource_id", ""),
                "class": loc.get("class", ""),
                "xpath": loc.get("xpath", ""),
                "uiautomator": _uiautomator(by, value, loc.get("resource_id", "")),
            },
        }
        by_identity[identity] = name
        names.append(name)
    return entries, names


def write_json(entries: Dict[str, Dict], path: Path, package: str) -> Path:
    data = {
        "format": "object-repository/1",
        "package": package,
        "generated": datetime.now().isoformat(timespec="seconds"),
        "selector_priority": ["text", "desc", "id", "xpath"],
        "count": len(entries),
        "elements": entries,
    }
    path.write_text(json.dumps(data, indent=2, ensure_ascii=False), encoding="utf-8")
    return path


def add_sheet(book, entries: Dict[str, Dict]) -> None:
    """The repository as a sheet of the MenuTree workbook."""
    from openpyxl.styles import Alignment, Font, PatternFill
    sheet = book.create_sheet("ObjectRepository")
    header = ["Name", "Screen", "Label", "Kind", "Preferred by", "Preferred value",
              "Text", "Content-desc", "Resource-id", "Class", "XPath", "UiAutomator"]
    for i, title in enumerate(header, start=1):
        cell = sheet.cell(row=1, column=i, value=title)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F3864")
    for r, e in enumerate(entries.values(), start=2):
        loc = e["locators"]
        values = [e["name"], e["screen"], e["label"], e["kind"],
                  e["preferred"]["by"], e["preferred"]["value"],
                  loc["text"], loc["content_desc"], loc["resource_id"], loc["class"],
                  loc["xpath"], loc["uiautomator"]]
        for c, v in enumerate(values, start=1):
            sheet.cell(row=r, column=c, value=v).alignment = Alignment(
                vertical="top", wrap_text=True)
    widths = [36, 30, 28, 10, 10, 28, 24, 24, 34, 24, 50, 44]
    from openpyxl.utils import get_column_letter
    for i, w in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(i)].width = w
    sheet.freeze_panes = "B2"
    sheet.auto_filter.ref = f"A1:L{max(1, len(entries) + 1)}"


def _identifier(text: str, taken: set) -> str:
    ident = _camel(text)
    if not ident or ident[0].isdigit():
        ident = "E" + ident
    if keyword.iskeyword(ident):
        ident += "_"
    base, n = ident, 2
    while ident in taken:
        ident, n = f"{base}_{n}", n + 1
    taken.add(ident)
    return ident


def write_page_objects(entries: Dict[str, Dict], path: Path, package: str) -> Path:
    """One class per screen, one Locator per element: page objects."""
    screens: "OrderedDict[str, List[Dict]]" = OrderedDict()
    for e in entries.values():
        screens.setdefault(e["screen"], []).append(e)
    out = [
        f'"""Page objects for {package} -- generated from a MenuTree run.',
        "",
        "Locators live here, apart from the tests. Each attribute is a Locator",
        "holding every way to find the element; `best` is the one the walk used",
        "(priority text > description > resource id > xpath).",
        "",
        "    from page_objects import PORTRAIT_QuickControls as Page",
        "    by, value = Page.Flash.best",
        '"""',
        "from typing import NamedTuple, Tuple",
        "",
        "",
        "class Locator(NamedTuple):",
        "    name: str",
        "    by: str",
        "    value: str",
        "    text: str = \"\"",
        "    content_desc: str = \"\"",
        "    resource_id: str = \"\"",
        "    class_name: str = \"\"",
        "    xpath: str = \"\"",
        "    uiautomator: str = \"\"",
        "    long_press: bool = False",
        "",
        "    @property",
        "    def best(self) -> Tuple[str, str]:",
        "        return (self.by, self.value)",
        "",
    ]
    class_names: set = set()
    for screen, items in screens.items():
        parts = [] if screen == "(main screen)" else screen.split(" > ")
        cls = "_".join(_segment(p) for p in parts) or "Main"
        if cls[0].isdigit():
            cls = "S" + cls
        base, n = cls, 2
        while cls in class_names:
            cls, n = f"{base}_{n}", n + 1
        class_names.add(cls)
        out += ["", f"class {cls}:", f"    SCREEN = {screen!r}"]
        attrs: set = set()
        for e in items:
            loc = e["locators"]
            attr = _identifier(_segment(e["label"]), attrs)
            out.append(
                f"    {attr} = Locator({e['name']!r}, {e['preferred']['by']!r}, "
                f"{e['preferred']['value']!r}, {loc['text']!r}, {loc['content_desc']!r}, "
                f"{loc['resource_id']!r}, {loc['class']!r}, {loc['xpath']!r}, "
                f"{loc['uiautomator']!r}, {e['long_press']!r})")
    path.write_text("\n".join(out) + "\n", encoding="utf-8")
    return path
