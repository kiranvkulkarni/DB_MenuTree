"""The Summary sheet: what a MenuTree was generated on (agreed 2026-10-20).

    python tests/test_run_info.py

No phone needed: adb is replaced by canned answers from a generic phone.
"""
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from src import run_info  # noqa: E402
from src.generator.menutree_workbook import write_workbook  # noqa: E402
from src.run_files import run_file  # noqa: E402

GETPROP = """[ro.product.manufacturer]: [examplecorp]
[ro.product.model]: [EX-100]
[ro.product.device]: [ex100]
[ro.build.version.release]: [16]
[ro.build.version.sdk]: [36]
[ro.build.version.oneui]: [80500]
[ro.build.version.security_patch]: [2026-06-05]
[ro.build.display.id]: [EX100-build 16 ABC.123]
[ro.build.PDA]: [EX100AP1]
[gsm.version.baseband]: [EX100CP1,EX100CP1]
[ro.revision]: [3]
[ro.soc.model]: [SOC9000]
[persist.sys.locale]: [en-US]
[ro.build.fingerprint]: [examplecorp/ex100/ex100:16/ABC.123/1:user/release-keys]
"""

DUMPSYS = """Packages:
  Package [com.example.app] (1234):
    codePath=/system/app/Example
    versionCode=1650238000 minSdk=29 targetSdk=36
    versionName=16.5.02.38
    lastUpdateTime=2026-01-02 03:04:05
"""


def fake_adb(serial, *args, timeout=20):
    joined = " ".join(args)
    if joined == "shell getprop":
        return GETPROP
    if joined.startswith("shell dumpsys package"):
        return DUMPSYS
    if "device_name" in joined:
        return "Example Phone\n"
    if joined == "shell wm size":
        return "Physical size: 1080x2340\n"
    if joined == "shell wm density":
        return "Physical density: 450\nOverride density: 480\n"
    return ""


def check(label, condition, detail=""):
    print(f"  [{'ok' if condition else 'FAIL'}] {label}"
          + (f"  ({detail})" if detail else ""))
    return condition


def main() -> int:
    ok = True
    run_info._adb = fake_adb

    print("phone and app read over adb")
    info = run_info.collect("com.example.app", "SER123", "Example")
    dev, app = info["device"], info["app"]
    ok &= check("phone name and model", dev.get("device_name") == "Example Phone"
                and dev.get("model") == "EX-100")
    ok &= check("software version", dev.get("android") == "16" and dev.get("sdk") == "36")
    ok &= check("vendor UI version made readable", dev.get("oneui") == "8.5 (80500)",
                dev.get("oneui"))
    ok &= check("build number, AP, CP", dev.get("build_number") == "EX100-build 16 ABC.123"
                and dev.get("ap") == "EX100AP1" and dev.get("cp") == "EX100CP1,EX100CP1")
    ok &= check("hardware revision and SoC", dev.get("hw_revision") == "3"
                and dev.get("soc") == "SOC9000")
    ok &= check("a property the phone lacks is left out", "csc" not in dev)
    ok &= check("density override wins", dev.get("screen_density") == "480")
    ok &= check("serial", dev.get("serial") == "SER123")
    ok &= check("app version", app.get("versionName") == "16.5.02.38"
                and app.get("versionCode") == "1650238000" and app.get("targetSdk") == "36")
    ok &= check("app name from the profile", info["app_name"] == "Example")

    print("no phone answers")
    run_info._adb = lambda *a, **k: ""
    empty = run_info.collect("com.example.app", None)
    ok &= check("empty sections, no error", empty["device"] == {} and empty["app"] == {})
    run_info._adb = fake_adb

    print("tree figures")
    rows = [
        {"label": "A", "depth": 1, "kind": "button", "path": []},
        {"label": "B", "depth": 2, "kind": "switch", "widget": "On", "path": ["A"]},
        {"label": "C", "depth": 3, "kind": "option", "path": ["A", "B"]},
        {"label": "C (long press)", "depth": 3, "kind": "button", "path": ["A", "B"]},
        {"label": "D", "depth": 1, "kind": "button", "path": []},
    ]
    figures = dict(run_info.tree_figures(rows))
    ok &= check("deepest level", figures["Deepest menu level (depth)"] == 3)
    ok &= check("top-level items", figures["Top-level items (depth 1)"] == 2)
    rooted = [{"label": "app", "depth": 1, "kind": "root", "path": []}] + [
        dict(r, depth=r["depth"] + 1) for r in rows]
    ok &= check("top-level items are under the root row",
                dict(run_info.tree_figures(rooted))["Top-level items (depth 2)"] == 2)
    ok &= check("rows per depth", figures["Rows at depth 3"] == 2)
    ok &= check("long presses counted", figures["Long-press rows"] == 1)

    print("workbook gets a Summary sheet, first")
    info["run"] = {"folder": "com.example.app_20261020_120000_000",
                   "started": "2026-10-20T12:00:00", "finished": "2026-10-20T14:05:00",
                   "complete": True, "time_budget": 28800, "reset": True}
    with tempfile.TemporaryDirectory() as tmp:
        out = run_file(Path(tmp) / "com.example.app_20261020_120000_000", "xlsx")
        write_workbook(rows, out, package="com.example.app",
                       stats={"elapsed_seconds": 7500, "clicks": 42},
                       run_info=info, extra_figures=[("UVTA test cases", 5)])
        from openpyxl import load_workbook
        book = load_workbook(out)
        ok &= check("sheet order", book.sheetnames[:2] == ["Summary", "MenuTree"],
                    book.sheetnames)
        sheet = book["Summary"]
        cells = {sheet.cell(r, 1).value: sheet.cell(r, 2).value
                 for r in range(1, sheet.max_row + 1)}
        ok &= check("headings", all(h in cells for h in (
            "MenuTree", "Application", "Phone", "Software version", "Build version",
            "Hardware version", "Tree", "Walk")))
        ok &= check("app name, version, package", cells.get("Application name") == "Example"
                    and cells.get("Version") == "16.5.02.38"
                    and cells.get("Package name") == "com.example.app")
        ok &= check("depth on the sheet", cells.get("Deepest menu level (depth)") == 3)
        ok &= check("walk duration", cells.get("Walk duration") == "2 h 05 min",
                    cells.get("Walk duration"))
        ok &= check("extra figures", cells.get("UVTA test cases") == 5)
        ok &= check("no run info, no Summary sheet", "Summary" not in load_workbook(
            write_workbook(rows, Path(tmp) / "b.xlsx", package="x")).sheetnames)

    print()
    print("ALL PASS" if ok else "FAILURES PRESENT")
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main())
